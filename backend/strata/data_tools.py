"""Pure, bounded file/data utilities. No solver, expression evaluation, DB or network.

All approval flags are caller-supplied declarations, not authenticated signatures.
Ingestion READY means structurally parsed, never engineering verification.
"""
from __future__ import annotations

import copy
import csv
import datetime as dt
import hashlib
import io
import json
import math
import re
import zipfile
from pathlib import PurePath
from typing import Any

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_CSV_BYTES = 1024 * 1024
MAX_ROWS = 5000
MAX_RECORDS = 1000
CSV_COLUMNS = ("record_type", "id", "floor_id", "case_id", "support_id", "value", "unit", "evidence_ref", "title", "locator", "content")
UNITS = {"m2": ("area", 1.0), "m²": ("area", 1.0), "kN": ("force", 1.0), "N": ("force", .001), "kN/m2": ("surfaceLoad", 1.0), "kN/m²": ("surfaceLoad", 1.0), "kPa": ("surfaceLoad", 1.0), "N/m2": ("surfaceLoad", .001), "N/m²": ("surfaceLoad", .001)}
CANONICAL = {"area": "m2", "force": "kN", "surfaceLoad": "kN/m2"}
NUMERIC = {"floor": ("floors", "area", "area", ["evidence_ref"]), "requirement": ("requirements", "q", "surfaceLoad", ["floor_id", "case_id", "evidence_ref"]), "assignment": ("assignments", "force", "force", ["floor_id", "case_id", "evidence_ref"]), "reaction": ("reactions", "fz", "force", ["support_id", "case_id", "evidence_ref"])}
FIELD_NAMES = {"floor_id": "floorId", "case_id": "caseId", "support_id": "supportId", "evidence_ref": "evidenceRef"}


class DataValidationError(ValueError):
    """Invalid or unsupported file/data; never silently repaired."""


def _require(condition: bool, message: str) -> None:
    if not condition:
        raise DataValidationError(message)


def _text(value: Any, maximum: int = 2000) -> bool:
    return isinstance(value, str) and bool(value.strip()) and len(value) <= maximum


def _number(value: Any) -> bool:
    try:
        return type(value) in (int, float) and math.isfinite(value)
    except OverflowError:
        return False


def _json(value: Any, depth: int = 0) -> None:
    _require(depth <= 32, "JSON nesting exceeds 32 levels.")
    if isinstance(value, str):
        try:
            value.encode("utf-8")
        except UnicodeEncodeError as error:
            raise DataValidationError("JSON strings must contain valid Unicode.") from error
        return
    if value is None or isinstance(value, bool):
        return
    if type(value) in (int, float):
        _require(_number(value), "JSON numbers must be finite.")
    elif isinstance(value, list):
        _require(len(value) <= 10000, "JSON collection too large.")
        for child in value:
            _json(child, depth + 1)
    elif isinstance(value, dict):
        _require(len(value) <= 10000 and all(isinstance(k, str) for k in value), "Invalid JSON object keys.")
        for key in value:
            _json(key, depth + 1)
        for child in value.values():
            _json(child, depth + 1)
    else:
        raise DataValidationError("Only finite JSON data is supported.")


def _hash(value: Any) -> str:
    _json(value)
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(",", ":"), allow_nan=False).encode()).hexdigest()


def _unique(values: list, label: str) -> None:
    _require(len(set(values)) == len(values), f"Duplicate {label}.")


def _schema(data: Any) -> None:
    _json(data)
    _require(isinstance(data, dict), "Input must be a JSON object.")
    _require(data.get("schemaVersion") in ("1.0", "combination-1.0"), "Unsupported schemaVersion.")
    _require(type(data.get("synthetic")) is bool, "synthetic must be an explicit boolean.")
    _require(isinstance(data.get("project"), dict) and _text(data["project"].get("name")), "project.name is required.")
    combination = data["schemaVersion"] == "combination-1.0"
    names = ("baseCases", "combinations", "evidence") if combination else ("floors", "requirements", "assignments", "reactions", "supports", "evidence")
    for name in names:
        _require(isinstance(data.get(name), list) and len(data[name]) <= MAX_RECORDS, f"{name} must be an array of at most {MAX_RECORDS} records.")
    for index, record in enumerate(data["evidence"]):
        _require(isinstance(record, dict) and _text(record.get("id")), f"evidence[{index}] needs an ID.")
        if combination:
            _require(all(record.get(k) is None or isinstance(record[k], str) for k in ("title", "locator", "content")), "Evidence text has an unsupported type.")
        else:
            _require(_text(record.get("title")) and all(isinstance(record.get(k), str) for k in ("locator", "content")), "Evidence needs a title and supplied text fields.")
    _unique([r["id"] for r in data["evidence"]], "evidence IDs")
    if combination:
        _require(_text(data["project"].get("revision")) and _text(data.get("unit")) and _text(data.get("analysisType")), "Combination revision, unit and analysisType are required.")
        for name, value_field in (("baseCases", "value"), ("combinations", "reportedValue")):
            for index, record in enumerate(data[name]):
                _require(isinstance(record, dict) and _text(record.get("id"), 160) and record["id"] == record["id"].strip() and _number(record.get(value_field)), f"{name}[{index}] has invalid ID or numeric value.")
                _require(record.get("evidenceRef") is None or isinstance(record["evidenceRef"], str), "Invalid evidenceRef type.")
                if name == "combinations":
                    terms = record.get("terms")
                    _require(isinstance(terms, list) and len(terms) <= MAX_RECORDS, "Combination terms must be a bounded array.")
                    for term in terms:
                        _require(isinstance(term, dict) and _text(term.get("caseId"), 160) and _number(term.get("factor")), "Invalid combination term.")
                    _unique([t["caseId"] for t in terms], "combination terms")
            _unique([r["id"] for r in data[name]], f"{name} IDs")
        _require(not ({r["id"] for r in data["baseCases"]} & {r["id"] for r in data["combinations"]}), "Base and combination IDs overlap.")
        return
    _require(isinstance(data.get("units"), dict) and isinstance(data.get("scope"), dict), "units and scope objects are required.")
    _require(bool(data["floors"]), "Floor list cannot be empty.")
    for record in data["floors"]:
        _require(isinstance(record, dict) and _text(record.get("id")) and _number(record.get("area")) and record["area"] > 0 and _text(record.get("evidenceRef")), "Invalid floor ID, area or evidenceRef.")
    _unique([r["id"] for r in data["floors"]], "floor IDs")
    _require(all(_text(s) for s in data["supports"]), "Invalid support ID.")
    _unique(data["supports"], "support IDs")
    for name, field, entity in (("requirements", "q", "floorId"), ("assignments", "force", "floorId"), ("reactions", "fz", "supportId")):
        for record in data[name]:
            _require(isinstance(record, dict) and all(_text(record.get(k)) for k in ("id", "caseId", "evidenceRef", entity)) and _number(record.get(field)) and record[field] >= 0, f"Invalid {name} record.")
    # Retain duplicate engineering records for the deterministic uniqueness check.
    areas = {r["id"]: r["area"] for r in data["floors"]}
    series = [[areas[r["floorId"]] * r["q"] for r in data["requirements"] if r["floorId"] in areas], [r["force"] for r in data["assignments"]], [r["fz"] for r in data["reactions"]]]
    for values in series:
        _require(all(_number(v) for v in values) and _number(sum(values)), "Engineering product or sum overflow.")


def _decode(content: bytes) -> str:
    try:
        text = content.decode("utf-8-sig", errors="strict")
    except UnicodeDecodeError as error:
        raise DataValidationError("Text input must be valid UTF-8.") from error
    _require("\x00" not in text, "NUL bytes are not allowed in text files.")
    return text


def _pairs(pairs: list) -> dict:
    result = {}
    for key, value in pairs:
        _require(key not in result, f"Duplicate JSON field: {key}.")
        result[key] = value
    return result


def _json_float(token: str) -> float:
    value = float(token)
    _require(_number(value), "JSON number overflow.")
    _require(not (value == 0 and re.search(r"[1-9]", re.split("[eE]", token)[0])), "JSON number underflow; value was not replaced with zero.")
    return value


def _pointer_token(value: Any) -> str:
    return str(value).replace("~", "~0").replace("/", "~1")


def _locations(data: Any, path: str = "") -> list:
    if isinstance(data, dict):
        return [item for key, value in data.items() for item in _locations(value, path + "/" + _pointer_token(key))]
    if isinstance(data, list):
        return [item for i, value in enumerate(data) for item in _locations(value, f"{path}/{i}")]
    return [{"path": path, "location": {"json_pointer": path}}]


def _controlled_csv(text: str, filename: str) -> tuple[dict, list]:
    # Python's csv strict mode still accepts quotes embedded in unquoted cells.
    # Reject those and characters after a closing quote before interpreting data.
    state = "start"
    for char in text:
        if state == "quoted":
            if char == '"': state = "closed"
        elif state == "closed":
            if char == '"': state = "quoted"
            elif char in ",\r\n": state = "start"
            else: raise DataValidationError("Invalid character after a CSV closing quote.")
        elif char == '"':
            _require(state == "start", "Quote inside an unquoted CSV cell.")
            state = "quoted"
        elif char in ",\r\n": state = "start"
        else: state = "unquoted"
    _require(state != "quoted", "Unterminated CSV quote.")
    reader = csv.reader(io.StringIO(text, newline=""), strict=True)
    try:
        header = next(reader)
        _require(len(header) == len(set(header)), "Duplicate CSV headers.")
        _require(set(header) == set(CSV_COLUMNS) and len(header) == len(CSV_COLUMNS), "CSV requires exactly the controlled column names.")
        data = {"project": {}, "units": {}, "scope": {}, **{name: [] for name in ("floors", "requirements", "assignments", "reactions", "supports", "evidence")}}
        metadata = {"format": "strata-controlled-csv", "version": "1.0", "filename": filename, "declaredUnits": {}, "rows": []}
        seen: set[tuple] = set()
        for cells in reader:
            # reader.line_num is the ending physical line; count embedded lines for start.
            row_num = reader.line_num - sum(len(re.findall(r"\r\n|\r|\n", value)) for value in cells)
            if not cells:
                continue
            _require(len(metadata["rows"]) < MAX_ROWS, "CSV row limit exceeded.")
            _require(len(cells) == len(header), f"CSV row {row_num}: incorrect column count.")
            row = dict(zip(header, cells))
            def required(key: str) -> str:
                _require(_text(row[key]), f"CSV row {row_num}, {key}: missing or oversized field.")
                return row[key]
            kind, rid = required("record_type"), required("id")
            numeric = {}
            if kind in NUMERIC:
                collection, field, dimension, fields = NUMERIC[kind]
                allowed = {"value", "unit", *fields}
                raw, unit = required("value"), required("unit")
                _require(bool(re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", raw.strip())), f"CSV row {row_num}: invalid decimal value.")
                value = float(raw)
                _require(_number(value) and value >= 0 and (field != "area" or value > 0), f"CSV row {row_num}: invalid numeric value.")
                _require(not (value == 0 and re.search(r"[1-9]", re.split("[eE]", raw)[0])), f"CSV row {row_num}: numeric underflow.")
                _require(unit in UNITS and UNITS[unit][0] == dimension, f"CSV row {row_num}: unknown or incompatible unit.")
                normalized = value * UNITS[unit][1]
                _require(_number(normalized) and (value == 0 or normalized != 0), f"CSV row {row_num}: unit conversion overflow or underflow.")
                record = {"id": rid, field: normalized, **{FIELD_NAMES[k]: required(k) for k in fields}}
                path = f"/{collection}/{len(data[collection])}"
                data[collection].append(record)
                numeric = {"field": field, "originalValue": raw, "originalUnit": unit, "normalizedValue": normalized, "normalizedUnit": CANONICAL[dimension]}
            elif kind in ("metadata", "project", "scope", "units"):
                allowed = {"unit"} if kind == "units" else {"value"}
                valid_ids = {"metadata": ["schemaVersion", "synthetic"], "project": ["name", "revision", "description", "software"], "scope": ["basis", "selfWeight", "reactionPositive"], "units": list(CANONICAL)}[kind]
                _require(rid in valid_ids, f"CSV row {row_num}: unknown {kind} field.")
                value = required("unit" if kind == "units" else "value")
                path = f"/{rid}" if kind == "metadata" else f"/{kind}/{rid}"
                if kind == "metadata":
                    _require(value == ("1.0" if rid == "schemaVersion" else "true"), f"CSV row {row_num}: unsupported schema or non-synthetic CSV.")
                    data[rid] = value if rid == "schemaVersion" else True
                elif kind == "units":
                    _require(value in UNITS and UNITS[value][0] == rid, f"CSV row {row_num}: unknown unit.")
                    data["units"][rid] = CANONICAL[rid]
                    metadata["declaredUnits"][rid] = value
                else:
                    data[kind][rid] = value
            elif kind == "support":
                allowed, path = set(), f"/supports/{len(data['supports'])}"
                data["supports"].append(rid)
            elif kind == "evidence":
                allowed, path = {"title", "locator", "content"}, f"/evidence/{len(data['evidence'])}"
                data["evidence"].append({"id": rid, "title": required("title"), "locator": row["locator"], "content": row["content"]})
            else:
                raise DataValidationError(f"CSV row {row_num}: unknown record type {kind}.")
            _require(all(not value for key, value in row.items() if key not in {"record_type", "id", *allowed}), f"CSV row {row_num}: populated unsupported cell.")
            if kind not in ("requirement", "assignment", "reaction"):
                _require((kind, rid) not in seen, f"CSV row {row_num}: duplicate {kind} ID.")
            seen.add((kind, rid))
            metadata["rows"].append({"row": row_num, "recordType": kind, "recordId": rid, "target": path, **numeric})
        required_records = [("metadata", "schemaVersion"), ("metadata", "synthetic"), ("project", "name"), *[("scope", k) for k in ("basis", "selfWeight", "reactionPositive")], *[("units", k) for k in CANONICAL]]
        _require(all(pair in seen for pair in required_records), "CSV missing a required metadata, scope or units record.")
        data["importMetadata"] = metadata
        _schema(data)
        return data, [{"path": row["target"], "location": {"row": row["row"]}, **row} for row in metadata["rows"]]
    except (csv.Error, StopIteration) as error:
        raise DataValidationError(f"Invalid CSV at physical line {reader.line_num}.") from error


def _xlsx(content: bytes) -> tuple[list, list, list]:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            _require(len(infos) <= 2000 and sum(i.file_size for i in infos) <= 50 * 1024 * 1024, "XLSX uncompressed size exceeds limit.")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except DataValidationError:
        raise
    except Exception as error:
        raise DataValidationError("Invalid or unsupported XLSX file.") from error
    tables, provenance, warnings = [], [], []
    total_cells = 0
    try:
        _require(len(workbook.worksheets) <= 20, "XLSX sheet limit exceeded.")
        for sheet in workbook.worksheets:
            _require(isinstance(sheet.max_row, int) and isinstance(sheet.max_column, int) and sheet.max_row <= MAX_ROWS and sheet.max_column <= 100, "XLSX dimensions missing or row/column limit exceeded.")
            table = {"name": sheet.title, "header_row": None, "columns": [], "rows": []}
            for cells in sheet.iter_rows():
                if all(cell.value is None for cell in cells):
                    continue
                row_number = next(cell.row for cell in cells if cell.value is not None)
                if table["header_row"] is None:
                    last = max(i for i, cell in enumerate(cells) if cell.value is not None)
                    header = cells[:last + 1]
                    _require(all(cell.data_type != "f" and _text(cell.value) for cell in header), f"{sheet.title}!{row_number}: headers must be nonempty text.")
                    table["columns"] = [cell.value.strip() for cell in header]
                    _unique(table["columns"], f"headers in {sheet.title}")
                    table["header_row"] = row_number
                    continue
                _require(all(cell.value is None for cell in cells[len(table['columns']):]), f"{sheet.title}!{row_number}: value outside named columns.")
                output_cells = []
                for index, cell in enumerate(cells[:len(table["columns"])]):
                    total_cells += 1
                    _require(total_cells <= 100000, "XLSX cell limit exceeded.")
                    value, cell_type = cell.value, "blank"
                    coordinate = getattr(cell, "coordinate", None) or f"{get_column_letter(index + 1)}{row_number}"
                    if cell.data_type == "f":
                        _require(isinstance(value, str), f"{sheet.title}!{coordinate}: unsupported formula type.")
                        cell_type = "formula"
                        warnings.append(f"{sheet.title}!{coordinate}: formula preserved; not evaluated.")
                    elif value is None:
                        pass
                    elif cell.data_type == "e":
                        raise DataValidationError(f"{sheet.title}!{coordinate}: spreadsheet error cell.")
                    elif type(value) is bool:
                        cell_type = "boolean"
                    elif _number(value):
                        cell_type = "number"
                    elif isinstance(value, str) and cell.data_type in ("s", "inlineStr", "str"):
                        cell_type = "text"
                    elif isinstance(value, (dt.datetime, dt.date, dt.time)):
                        cell_type, value = "date", value.isoformat()
                    else:
                        raise DataValidationError(f"{sheet.title}!{coordinate}: unsupported cell type.")
                    output_cells.append({"column": table["columns"][index], "coordinate": coordinate, "type": cell_type, "value": value})
                    provenance.append({"path": f"/tables/{len(tables)}/rows/{len(table['rows'])}/cells/{index}", "location": {"sheet": sheet.title, "row": row_number, "cell": coordinate}})
                table["rows"].append({"row": row_number, "cells": output_cells})
            tables.append(table)
    except DataValidationError:
        raise
    except Exception as error:
        raise DataValidationError("Invalid or unsupported XLSX cell data.") from error
    finally:
        workbook.close()
    return tables, provenance, warnings


def ingest(filename: str, content: bytes) -> dict:
    """Parse an uploaded file. Errors raise DataValidationError; no files are written."""
    _require(_text(filename, 255) and not any(c in filename for c in ("/", "\\", "\x00")), "Provide a plain filename, not a path.")
    _require(type(content) is bytes and 0 < len(content) <= MAX_FILE_BYTES, "File must contain 1 byte to 10 MiB.")
    extension = PurePath(filename).suffix.lower()
    _require(extension in (".json", ".csv", ".xlsx", ".pdf", ".txt", ".md"), "Unsupported file extension.")
    result = {"status": "READY", "format": extension[1:], "file": {"name": filename, "sha256": hashlib.sha256(content).hexdigest(), "size_bytes": len(content)}, "snapshot": None, "tables": [], "pages": [], "provenance": [], "warnings": []}
    if extension == ".json":
        try:
            raw_text = _decode(content)
            data = json.loads(raw_text, object_pairs_hook=_pairs, parse_float=_json_float, parse_constant=lambda value: (_ for _ in ()).throw(DataValidationError("Non-finite JSON number.")))
        except (json.JSONDecodeError, RecursionError) as error:
            raise DataValidationError("Malformed JSON input.") from error
        _json(data)
        if isinstance(data, dict) and data.get("schemaVersion") in ("1.0", "combination-1.0"):
            _schema(data)
            result.update(snapshot=data, provenance=_locations(data))
        else:
            result.update(status="NEEDS_MAPPING", pages=[{"page": 1, "text": raw_text}], provenance=[{"path": "/pages/0/text", "location": {"page": 1}}])
            result["warnings"].append("JSON document has no supported engineering schema; no engineering snapshot was inferred.")
    elif extension in (".txt", ".md"):
        result.update(pages=[{"page": 1, "text": _decode(content)}], provenance=[{"path": "/pages/0/text", "location": {"page": 1}}])
        result["warnings"].append("Text is a source document only; no engineering snapshot or rule approval was inferred.")
    elif extension == ".csv":
        _require(len(content) <= MAX_CSV_BYTES, "Controlled CSV exceeds 1 MiB.")
        result["snapshot"], result["provenance"] = _controlled_csv(_decode(content), filename)
        result["format"] = "controlled-csv"
    elif extension == ".xlsx":
        result["tables"], result["provenance"], result["warnings"] = _xlsx(content)
        result["status"] = "NEEDS_MAPPING"
        result["warnings"].append("Workbook cells require an explicit engineering mapping; formulas were not calculated.")
    else:
        from pypdf import PdfReader
        _require(content.startswith(b"%PDF-"), "PDF signature does not match extension.")
        try:
            reader = PdfReader(io.BytesIO(content), strict=True)
            _require(not reader.is_encrypted, "Encrypted PDF is unsupported.")
            _require(0 < len(reader.pages) <= 200, "PDF must have 1 to 200 pages.")
            total = 0
            for number, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                total += len(text)
                _require(len(text) <= 100000 and total <= 2000000, "PDF extracted text exceeds limit.")
                result["pages"].append({"page": number, "text": text, "needs_ocr": not bool(text.strip())})
                result["provenance"].append({"path": f"/pages/{number - 1}/text", "location": {"page": number}})
        except DataValidationError:
            raise
        except Exception as error:
            raise DataValidationError("PDF text extraction failed; no content was inferred.") from error
        needs_ocr = any(p["needs_ocr"] for p in result["pages"])
        result["status"] = "NEEDS_OCR" if needs_ocr else "NEEDS_MAPPING"
        result["warnings"].append("Pages without extractable text may be scanned or blank; OCR/manual review is required. No OCR was performed." if needs_ocr else "Extracted text requires an explicit engineering mapping; no formulas were inferred.")
    return result


GRAVITY_DEPENDENCIES = {"floors": ["QA-002", "QA-003", "QA-004", "QA-005", "QA-006"], "requirements": ["QA-001", "QA-002", "QA-003", "QA-004", "QA-005", "QA-006"], "assignments": ["QA-001", "QA-002", "QA-003", "QA-004", "QA-006"], "reactions": ["QA-001", "QA-005", "QA-006"], "supports": ["QA-005"], "evidence": ["QA-002", "QA-003", "QA-004", "QA-005", "QA-006"], "units": ["QA-003", "QA-004", "QA-005"], "scope": ["QA-003", "QA-004", "QA-005"]}
ALL_GRAVITY = [f"QA-00{i}" for i in range(1, 7)]


def compare_snapshots(before: dict, after: dict) -> dict:
    """Compare data by stable IDs where unique; entity paths are display paths, not JSON pointers."""
    _json(before); _json(after)
    _require(isinstance(before, dict) and isinstance(after, dict), "Snapshots must be JSON objects.")
    changes, unknown_reasons = [], []
    def add(path, entity, kind, left, right):
        changes.append({"path": path or "/", "entity": entity, "change": kind, "before": copy.deepcopy(left), "after": copy.deepcopy(right)})
    def walk(left, right, path="", entity=None):
        # Python equates True with 1, including inside collections; JSON does not.
        if json.dumps(left, sort_keys=True, ensure_ascii=False) == json.dumps(right, sort_keys=True, ensure_ascii=False):
            return
        if isinstance(left, dict) and isinstance(right, dict):
            for key in sorted(set(left) | set(right)):
                location = path + "/" + _pointer_token(key)
                if key not in left: add(location, entity, "added", None, right[key])
                elif key not in right: add(location, entity, "removed", left[key], None)
                else: walk(left[key], right[key], location, entity)
        elif isinstance(left, list) and isinstance(right, list):
            records = left + right
            identifiable = all(isinstance(r, dict) and _text(r.get("id")) for r in records)
            if identifiable and len({r["id"] for r in left}) == len(left) and len({r["id"] for r in right}) == len(right):
                a, b = {r["id"]: r for r in left}, {r["id"]: r for r in right}
                for rid in sorted(set(a) | set(b)):
                    location, record_entity = path + "/@" + _pointer_token(rid), path.strip("/") + ":" + rid
                    if rid not in a: add(location, record_entity, "added", None, b[rid])
                    elif rid not in b: add(location, record_entity, "removed", a[rid], None)
                    else: walk(a[rid], b[rid], location, record_entity)
            else:
                if identifiable: unknown_reasons.append(f"Ambiguous duplicate entity IDs at {path}; compare full sequence.")
                add(path, entity, "modified", left, right)
        else:
            add(path, entity, "modified", left, right)
    walk(before, after)
    checks = set()
    schema = after.get("schemaVersion")
    if schema != before.get("schemaVersion") or schema not in ("1.0", "combination-1.0"):
        unknown_reasons.append("Schema changed or is unsupported; full review required.")
    for change in changes:
        root = change["path"].split("/")[1]
        if root in ("project", "synthetic", "schemaVersion"):
            checks.update(ALL_GRAVITY if schema == "1.0" else ["COMB-001", "COMB-CONFIG"])
        elif schema == "1.0" and root in GRAVITY_DEPENDENCIES:
            checks.update(GRAVITY_DEPENDENCIES[root])
        elif schema == "combination-1.0" and root in ("baseCases", "combinations", "evidence", "analysisType", "unit", "configuration_complete"):
            checks.update(["COMB-001", "COMB-CONFIG"])
        else:
            unknown_reasons.append(f"No complete dependency definition for {root}.")
    if unknown_reasons: checks.add("ALL_SUPPORTED_CHECKS")
    return {"status": "CHANGED" if changes else "UNCHANGED", "before_hash": _hash(before), "after_hash": _hash(after), "changes": changes, "summary": {kind: sum(c["change"] == kind for c in changes) for kind in ("added", "removed", "modified")}, "impacts": {"checks": sorted(checks), "unknown": bool(unknown_reasons), "reasons": sorted(set(unknown_reasons))}}


def _pointer(data: Any, path: str) -> Any:
    _require(isinstance(path, str) and path.startswith("/") and not re.search(r"~(?![01])", path), "Mapping paths must be JSON pointers.")
    node = data
    for token in path[1:].split("/"):
        token = token.replace("~1", "/").replace("~0", "~")
        if isinstance(node, list):
            _require(bool(re.fullmatch(r"0|[1-9]\d*", token)) and int(token) < len(node), f"Missing mapped path {path}.")
            node = node[int(token)]
        else:
            _require(isinstance(node, dict) and token in node, f"Missing mapped path {path}.")
            node = node[token]
    return node


def _outcome(details: list, scope: str, reason: str = "") -> dict:
    status = "FAIL" if any(d["status"] == "FAIL" for d in details) else "NOT VERIFIED" if not details or any(d["status"] == "NOT VERIFIED" for d in details) else "PASS"
    return {"status": status, "reason": reason or ("All checks within the explicitly declared scope passed." if status == "PASS" else "See per-item findings and missing prerequisites."), "scope": scope, "details": details, "dependencies": sorted({p for d in details for p in d.get("dependencies", [])})}


def verify_handoff(source: dict, target: dict, manifest: dict) -> dict:
    """Verify only explicit one-to-one scalar mappings. No solver or inferred mappings."""
    scope = "Only the complete field set explicitly listed in the supplied approved manifest; not model safety or native CSI integration."
    try:
        _json(source); _json(target); _json(manifest)
        required = {"id", "version", "approved", "source_revision", "target_revision", "required_source_paths", "required_target_paths", "mappings"}
        _require(isinstance(manifest, dict) and set(manifest) == required and manifest["approved"] is True and all(_text(manifest[k]) for k in ("id", "version", "source_revision", "target_revision")), "A complete, explicitly approved versioned manifest is required.")
        _require(source.get("project", {}).get("revision") == manifest["source_revision"] and target.get("project", {}).get("revision") == manifest["target_revision"], "Snapshot revisions do not match the manifest.")
        for key in ("required_source_paths", "required_target_paths"):
            _require(isinstance(manifest[key], list) and 0 < len(manifest[key]) <= MAX_RECORDS and all(_text(p) and p.startswith("/") for p in manifest[key]), "Required mapped path sets must be nonempty.")
            _unique(manifest[key], key)
        mappings = manifest["mappings"]
        _require(isinstance(mappings, list) and 0 < len(mappings) <= MAX_RECORDS, "No approved mappings supplied.")
        fields = {"id", "source_path", "target_path", "source_unit_path", "target_unit_path", "comparison_unit", "absolute_tolerance", "relative_tolerance"}
        for mapping in mappings:
            _require(isinstance(mapping, dict) and set(mapping) == fields and all(_text(mapping[k]) for k in fields - {"absolute_tolerance", "relative_tolerance"}), "Incomplete or unsupported mapping definition.")
            _require(all(_number(mapping[k]) and mapping[k] >= 0 for k in ("absolute_tolerance", "relative_tolerance")), "Explicit finite nonnegative tolerances are required.")
        for field in ("id", "source_path", "target_path"):
            _unique([m[field] for m in mappings], f"mapping {field}")
        _require(set(manifest["required_source_paths"]) == {m["source_path"] for m in mappings} and set(manifest["required_target_paths"]) == {m["target_path"] for m in mappings}, "Required fields are missing from the approved mappings.")
    except (DataValidationError, AttributeError, TypeError) as error:
        return _outcome([], scope, str(error))
    details = []
    for mapping in mappings:
        detail = {"id": mapping["id"], "status": "NOT VERIFIED", "reason": "", "dependencies": [mapping["source_path"], mapping["target_path"], mapping["source_unit_path"], mapping["target_unit_path"]]}
        try:
            original_source, original_target = _pointer(source, mapping["source_path"]), _pointer(target, mapping["target_path"])
            source_unit, target_unit = _pointer(source, mapping["source_unit_path"]), _pointer(target, mapping["target_unit_path"])
            unit = mapping["comparison_unit"]
            _require(_number(original_source) and _number(original_target), "Mapped values must be finite numbers; missing values are not zero.")
            _require(all(isinstance(u, str) and u in UNITS for u in (source_unit, target_unit, unit)), "Unknown unit; no conversion inferred.")
            _require(UNITS[source_unit][0] == UNITS[target_unit][0] == UNITS[unit][0], "Incompatible unit dimensions.")
            expected = original_source * (UNITS[source_unit][1] / UNITS[unit][1])
            actual = original_target * (UNITS[target_unit][1] / UNITS[unit][1])
            tolerance = max(mapping["absolute_tolerance"], abs(expected) * mapping["relative_tolerance"])
            delta = actual - expected
            _require(all(_number(v) for v in (expected, actual, tolerance, delta)), "Handoff arithmetic overflow.")
            _require(not ((original_source != 0 and expected == 0) or (original_target != 0 and actual == 0)), "Handoff unit conversion underflow.")
            detail.update(status="PASS" if abs(delta) <= tolerance else "FAIL", source_value=original_source, target_value=original_target, source_unit=source_unit, target_unit=target_unit, expected=expected, actual=actual, unit=unit, delta=delta, tolerance=tolerance, reason="Compared against explicit manifest tolerance.")
        except DataValidationError as error:
            detail["reason"] = str(error)
        details.append(detail)
    return {**_outcome(details, scope), "manifest": {"id": manifest["id"], "version": manifest["version"]}}


def verify_combination_configuration(input: dict, rule: dict) -> dict:
    """Check declared combination membership/factors, independently of response arithmetic.

    input uses baseCases/combinations and configuration_complete. Terms are exact sets
    for each required combination. Approval is a supplied declaration, not certified.
    """
    scope = "Supplied approved combination configuration only; no interpretation of structural codes or response calculation."
    try:
        _json(input); _json(rule)
        required = {"id", "version", "approved", "required_base_cases", "required_combinations", "allow_extra_base_cases", "allow_extra_combinations", "factor_tolerance"}
        _require(isinstance(rule, dict) and set(rule) == required and rule["approved"] is True and _text(rule["id"]) and _text(rule["version"]), "An explicitly approved versioned configuration rule is required.")
        _require(type(rule["allow_extra_base_cases"]) is bool and type(rule["allow_extra_combinations"]) is bool and _number(rule["factor_tolerance"]) and rule["factor_tolerance"] >= 0, "Explicit extra-item policies and factor tolerance are required.")
        _require(isinstance(rule["required_base_cases"], list) and 0 < len(rule["required_base_cases"]) <= MAX_RECORDS and all(_text(v) for v in rule["required_base_cases"]), "Required base cases must be explicit and bounded.")
        _unique(rule["required_base_cases"], "required base cases")
        _require(isinstance(rule["required_combinations"], list) and 0 < len(rule["required_combinations"]) <= MAX_RECORDS, "Required combinations must be explicit and bounded.")
        for record in rule["required_combinations"]:
            _require(isinstance(record, dict) and set(record) == {"id", "terms"} and _text(record["id"]) and isinstance(record["terms"], list) and 0 < len(record["terms"]) <= MAX_RECORDS, "Invalid required combination.")
            for term in record["terms"]:
                _require(isinstance(term, dict) and set(term) == {"caseId", "factor"} and term["caseId"] in rule["required_base_cases"] and _number(term["factor"]), "Required factors must be finite and reference a declared required base case.")
            _unique([t["caseId"] for t in record["terms"]], "required terms")
        _unique([r["id"] for r in rule["required_combinations"]], "required combinations")
        _require(not (set(rule["required_base_cases"]) & {r["id"] for r in rule["required_combinations"]}), "Required base case and combination IDs overlap.")
        _require(isinstance(input, dict) and input.get("schemaVersion") == "combination-1.0", "Expected a combination-1.0 configuration snapshot.")
        for name in ("baseCases", "combinations"):
            _require(isinstance(input.get(name), list) and len(input[name]) <= MAX_RECORDS, f"Missing or oversized {name} list.")
            _require(all(isinstance(r, dict) and _text(r.get("id")) for r in input[name]), f"Malformed {name} entry.")
        for record in input["combinations"]:
            _require(isinstance(record.get("terms"), list) and len(record["terms"]) <= MAX_RECORDS, "Missing combination terms.")
            _require(all(isinstance(t, dict) and _text(t.get("caseId")) and _number(t.get("factor")) for t in record["terms"]), "Malformed configuration term.")
    except (DataValidationError, TypeError, KeyError) as error:
        return _outcome([], scope, str(error))
    details = []
    complete = input.get("configuration_complete") is True
    def record(identifier, status, reason, dependencies):
        details.append({"id": identifier, "status": status, "reason": reason, "dependencies": dependencies})
    if not complete:
        record("completeness", "NOT VERIFIED", "Export completeness is not explicitly declared; absence cannot establish a violation.", ["/configuration_complete"])
    if {r["id"] for r in input["baseCases"]} & {r["id"] for r in input["combinations"]}:
        record("namespace", "FAIL", "Base case and combination IDs overlap; no nested interpretation was inferred.", ["/baseCases", "/combinations"])
    for name, expected, allow_extra in (("baseCases", rule["required_base_cases"], rule["allow_extra_base_cases"]), ("combinations", [r["id"] for r in rule["required_combinations"]], rule["allow_extra_combinations"])):
        ids = [r["id"] for r in input[name]]
        if len(ids) != len(set(ids)):
            record(name + ":duplicates", "FAIL", "Duplicate configuration IDs; records were not merged.", ["/" + name])
        missing, extra = sorted(set(expected) - set(ids)), sorted(set(ids) - set(expected))
        status = "FAIL" if (extra and not allow_extra) or (missing and complete) else "NOT VERIFIED" if missing or not complete else "PASS"
        record(name + ":membership", status, f"Missing: {missing}; extra: {extra}; extra allowed: {allow_extra}.", ["/" + name, "/configuration_complete"])
    for required_combination in rule["required_combinations"]:
        matches = [r for r in input["combinations"] if r["id"] == required_combination["id"]]
        if len(matches) != 1:
            continue
        combination = matches[0]
        terms = combination["terms"]
        ids = [t["caseId"] for t in terms]
        expected = {t["caseId"]: t["factor"] for t in required_combination["terms"]}
        if len(ids) != len(set(ids)):
            record(combination["id"] + ":terms", "FAIL", "Duplicate term IDs; factors were not combined.", ["/combinations"])
            continue
        actual = {t["caseId"]: t["factor"] for t in terms}
        missing, extra = sorted(set(expected) - set(actual)), sorted(set(actual) - set(expected))
        wrong = [key for key in set(actual) & set(expected) if abs(actual[key] - expected[key]) > rule["factor_tolerance"]]
        status = "FAIL" if wrong or extra or (missing and complete) else "NOT VERIFIED" if missing or not complete else "PASS"
        record(combination["id"] + ":terms", status, f"Missing terms: {missing}; extra terms: {extra}; factors outside tolerance: {sorted(wrong)}.", ["/combinations", "/configuration_complete"])
    return {**_outcome(details, scope), "rule": {"id": rule["id"], "version": rule["version"]}}


def verify_settings(input: dict, rule: dict) -> dict:
    """Compare explicit configuration settings only; no structural-code inference.

    Rule shape: {id, version, approved, settings: [{path, operator, value} |
    {path, operator: 'range', min, max}]}. Range bounds are inclusive. `eq` and
    `in` support non-null scalar JSON values; boolean and number are distinct.
    An application must supply approval from its review store, not trust uploads.
    """
    scope = "Only the settings and bounds explicitly declared by the supplied approved rule; not structural safety or code compliance."
    try:
        _json(input); _json(rule)
        _require(isinstance(input, dict), "Configuration input must be an object.")
        _require(isinstance(rule, dict) and set(rule) == {"id", "version", "approved", "settings"} and rule["approved"] is True and _text(rule["id"]) and _text(rule["version"]), "An explicitly approved versioned settings rule is required.")
        settings = rule["settings"]
        _require(isinstance(settings, list) and 0 < len(settings) <= MAX_RECORDS, "An explicit bounded nonempty settings list is required.")
    except DataValidationError as error:
        return _outcome([], scope, str(error))
    details = []
    if input.get("configuration_complete") is not True:
        details.append({"id": "completeness", "status": "NOT VERIFIED", "reason": "Configuration completeness must be explicitly declared.", "dependencies": ["/configuration_complete"]})

    def scalar(value):
        return isinstance(value, (str, bool)) or _number(value)

    def equal(left, right):
        if _number(left) and _number(right):
            return left == right
        return type(left) is type(right) and left == right

    for index, setting in enumerate(settings):
        detail = {"id": f"setting-{index + 1}", "status": "NOT VERIFIED", "reason": "", "dependencies": []}
        try:
            _require(isinstance(setting, dict) and _text(setting.get("path")), "Each setting requires a JSON pointer path.")
            path, operator = setting["path"], setting.get("operator")
            detail["path"] = path
            detail["dependencies"] = [path, "/configuration_complete"]
            _require(operator in ("eq", "in", "range"), "Unsupported setting operator; no comparison inferred.")
            expected_keys = {"path", "operator", "min", "max"} if operator == "range" else {"path", "operator", "value"}
            _require(set(setting) == expected_keys, "Incomplete or unsupported setting fields.")
            if operator == "range":
                _require(_number(setting["min"]) and _number(setting["max"]) and setting["min"] <= setting["max"], "Range requires ordered finite numeric bounds.")
            elif operator == "eq":
                _require(scalar(setting["value"]), "Equality requires an explicit non-null scalar value.")
            else:
                _require(isinstance(setting["value"], list) and 0 < len(setting["value"]) <= MAX_RECORDS and all(scalar(v) for v in setting["value"]), "Membership requires a nonempty bounded list of non-null scalar values.")
            actual = _pointer(input, path)
            _require(scalar(actual), "Setting has a missing, null or unsupported value; it was not replaced with a default.")
            if operator == "range":
                _require(_number(actual), "Numeric range cannot compare boolean or non-numeric input.")
                passed = setting["min"] <= actual <= setting["max"]
            elif operator == "eq":
                passed = equal(actual, setting["value"])
            else:
                passed = any(equal(actual, value) for value in setting["value"])
            detail.update(status="PASS" if passed else "FAIL", operator=operator, actual=actual, reason="Compared against the explicit approved setting requirement.")
            detail.update({key: copy.deepcopy(value) for key, value in setting.items() if key in ("value", "min", "max")})
        except DataValidationError as error:
            detail["reason"] = str(error)
        details.append(detail)
    return {**_outcome(details, scope), "rule": {"id": rule["id"], "version": rule["version"]}}
